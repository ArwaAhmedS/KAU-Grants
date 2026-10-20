"""
form_delivery.py — إرسال نماذج الموقع العام بالإيميل، وحفظ طلبات تقديم المنح في ملف Excel.
"""
import os

from django.conf import settings
from django.core.mail import EmailMultiAlternatives
from django.template.loader import render_to_string
from openpyxl import Workbook, load_workbook


# ملف Excel التراكمي لطلبات تقديم المنح (داخل مجلد media/exports)
GRANT_EXCEL_PATH = os.path.join(settings.MEDIA_ROOT, "exports", "grant_requests.xlsx")


def send_request_email(subject, rows):
    """يرسل إيميلاً موحّداً (نص + HTML) لإدارة علاقات المانحين.
    rows: قائمة عناصر بالشكل {"label": ..., "value": ...} تُعرض كجدول داخل الإيميل."""
    plain_body = "\n".join(f"{row['label']}: {row['value']}" for row in rows)
    html_body = render_to_string("public_site/emails/request_email.html", {"title": subject, "rows": rows})

    msg = EmailMultiAlternatives(
        subject=subject,
        body=plain_body,
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=[settings.GRANT_REQUESTS_RECIPIENT],
    )
    msg.attach_alternative(html_body, "text/html")
    msg.send(fail_silently=False)


def append_grant_to_excel(row_dict):
    """يضيف صفاً جديداً لملف طلبات المنح (وينشئ الملف مع صف العناوين لو لم يكن موجوداً)."""
    os.makedirs(os.path.dirname(GRANT_EXCEL_PATH), exist_ok=True)
    if os.path.exists(GRANT_EXCEL_PATH):
        wb = load_workbook(GRANT_EXCEL_PATH)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "طلبات تقديم منح"
        ws.append(list(row_dict.keys()))
    ws.append(list(row_dict.values()))
    # القيم قادمة من زوار الموقع: نص يبدأ بـ "=" يُحفظ كنص عادي وليس كمعادلة،
    # حتى لا تُنفَّذ معادلة ضارة عند فتح الملف في Excel
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"
    wb.save(GRANT_EXCEL_PATH)
